from __future__ import annotations

import argparse
import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(message)s")
log = logging.getLogger("task4")

TEXT_COLS = [
    "Transaction ID", "Customer ID", "Category", "Item",
    "Payment Method", "Location"
]
MODEL_COLS = [
    "Transaction ID", "Customer ID", "Category", "Item",
    "Price Per Unit", "Quantity", "Total Spent", "Payment Method",
    "Location", "Transaction Date", "Discount Applied"
]
NULL_TOKENS = ["", "nan", "NaN", "NULL", "null", "None", "ERROR", "UNKNOWN", "Unknown"]


def clean_data(raw: pd.DataFrame):
    df = raw.copy()
    actions = []

    def log_action(step, rows, note):
        actions.append((step, int(rows), note))
        log.info("%-34s %6d rows | %s", step, int(rows), note)

    # Standardise text and null-like tokens.
    for col in TEXT_COLS:
        df[col] = df[col].astype("string").str.strip().replace(NULL_TOKENS, pd.NA)
    log_action("Standardise text", len(df), "Stripped whitespace and standardised null-like tokens")

    # Duplicate handling is defensive: this dataset has none, but the automation handles them.
    n = int(df.duplicated().sum())
    df = df.drop_duplicates()
    log_action("Remove duplicate rows", n, "Exact duplicate rows dropped")

    n = int(df["Transaction ID"].duplicated().sum())
    df = df.drop_duplicates(subset="Transaction ID", keep="first")
    log_action("Remove duplicate Transaction IDs", n, "Kept first occurrence")

    # Types.
    df["Transaction Date"] = pd.to_datetime(df["Transaction Date"], errors="coerce")
    for col in ["Price Per Unit", "Quantity", "Total Spent"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    log_action("Parse dates / numerics", int(df["Transaction Date"].isna().sum()), "Invalid dates become NaT")

    # Discount: preserve uncertainty instead of guessing.
    n = int(df["Discount Applied"].isna().sum())
    df["Discount Applied"] = (
        df["Discount Applied"].astype("string").str.strip().str.lower()
        .map({"true": "Yes", "false": "No"}).fillna("Unknown")
    )
    log_action("Discount Applied", n, "Missing flags labelled Unknown; no assumption made")

    # Audit flags are kept separately from the final 11-column clean dataset.
    df["Price_Recovered"] = False
    df["Item_Recovered"] = False
    df["Quantity_Imputed"] = False
    df["Total_Recalculated"] = False

    # Recover Item only where Category + Price uniquely identifies one Item.
    lookup = (
        df.dropna(subset=["Item", "Price Per Unit"])
        .groupby(["Category", "Price Per Unit"])["Item"]
        .agg(lambda s: s.iloc[0] if s.nunique() == 1 else pd.NA)
        .dropna()
    )
    miss = df["Item"].isna() & df["Price Per Unit"].notna()
    if miss.any():
        keys = pd.MultiIndex.from_frame(df.loc[miss, ["Category", "Price Per Unit"]])
        recovered = pd.Series(lookup.reindex(keys).to_numpy(), index=df.index[miss])
        valid = recovered.notna()
        df.loc[recovered.index[valid], "Item"] = recovered[valid]
        df.loc[recovered.index[valid], "Item_Recovered"] = True
        recovered_count = int(valid.sum())
    else:
        recovered_count = 0
    log_action("Recover Item from price", recovered_count, "Only unique Category + Price mappings used")

    # Remaining unknown Item values are explicit, not fabricated.
    n = int(df["Item"].isna().sum())
    df["Item"] = df["Item"].fillna("Unknown Item")
    log_action("Unknown Item label", n, "Could not be safely recovered")

    # Recover Price from known Item first.
    item_price = (
        df.dropna(subset=["Price Per Unit"])
        .loc[df["Item"] != "Unknown Item"]
        .groupby("Item")["Price Per Unit"]
        .agg(lambda s: s.mode().iloc[0])
    )
    miss = df["Price Per Unit"].isna() & df["Item"].isin(item_price.index)
    df.loc[miss, "Price Per Unit"] = df.loc[miss, "Item"].map(item_price)
    df.loc[miss, "Price_Recovered"] = True
    log_action("Recover Price from item", int(miss.sum()), "Used the modal known price for the item")

    # Then use the proven Total / Quantity relationship.
    miss = (
        df["Price Per Unit"].isna()
        & df["Quantity"].notna()
        & df["Total Spent"].notna()
        & (df["Quantity"] != 0)
    )
    df.loc[miss, "Price Per Unit"] = df.loc[miss, "Total Spent"] / df.loc[miss, "Quantity"]
    df.loc[miss, "Price_Recovered"] = True
    log_action("Recover Price from Total/Qty", int(miss.sum()), "Total Spent / Quantity")

    # Recover Quantity exactly when Price and Total are available.
    miss = (
        df["Quantity"].isna()
        & df["Total Spent"].notna()
        & df["Price Per Unit"].notna()
        & (df["Price Per Unit"] != 0)
    )
    recovered_qty = (df.loc[miss, "Total Spent"] / df.loc[miss, "Price Per Unit"]).round()
    df.loc[miss, "Quantity"] = recovered_qty
    log_action("Recover Quantity from Total/Price", int(miss.sum()), "Exact recovery where mathematically determined")

    # Remaining Quantity values use category median and are explicitly flagged.
    med = df.groupby("Category")["Quantity"].transform("median").round()
    miss = df["Quantity"].isna()
    df.loc[miss, "Quantity"] = med[miss]
    df.loc[miss, "Quantity_Imputed"] = True
    log_action("Impute Quantity (median)", int(miss.sum()), "Category median; flagged as estimated")

    # Recalculate missing totals.
    miss = df["Total Spent"].isna() & df["Price Per Unit"].notna() & df["Quantity"].notna()
    df.loc[miss, "Total Spent"] = df.loc[miss, "Price Per Unit"] * df.loc[miss, "Quantity"]
    df.loc[miss, "Total_Recalculated"] = True
    log_action("Recalculate Total Spent", int(miss.sum()), "Price Per Unit x Quantity")

    df["Is_Estimated"] = df["Quantity_Imputed"]
    return df.reset_index(drop=True), actions


def validate(raw: pd.DataFrame, clean: pd.DataFrame) -> pd.DataFrame:
    model = clean[MODEL_COLS]
    diff = (clean["Total Spent"] - clean["Price Per Unit"] * clean["Quantity"]).abs()
    rows = [
        ("Rows (raw)", len(raw), ""),
        ("Rows (clean)", len(model), ""),
        ("Missing cells (raw)", int(raw.isna().sum().sum()), ""),
        ("Missing cells (clean)", int(model.isna().sum().sum()), "Target = 0"),
        ("Duplicate rows (clean)", int(model.duplicated().sum()), "Target = 0"),
        ("Duplicate Transaction IDs (clean)", int(model["Transaction ID"].duplicated().sum()), "Target = 0"),
        ("Invalid dates (clean)", int(model["Transaction Date"].isna().sum()), "Target = 0"),
        ("Negative price/qty/total", int((model[["Price Per Unit", "Quantity", "Total Spent"]] < 0).any(axis=1).sum()), "Target = 0"),
        ("Total != Price x Qty", int((diff > 0.01).sum()), "Target = 0"),
        ("Estimated rows (Quantity imputed)", int(clean["Is_Estimated"].sum()), "Flagged, not hidden"),
    ]
    out = pd.DataFrame(rows, columns=["Check", "Result", "Note"])
    print("\nVALIDATION RESULTS\n" + "-" * 18)
    print(out.to_string(index=False))
    return out


def analyse(df: pd.DataFrame) -> dict:
    verified = df.loc[~df["Is_Estimated"], "Total Spent"].sum()
    total = df["Total Spent"].sum()
    kpi = pd.DataFrame({
        "KPI": [
            "Total Revenue (incl. estimated)",
            "Verified Revenue (excl. estimated rows)",
            "Estimated Revenue Share",
            "Total Transactions",
            "Unique Customers",
            "Average Transaction Value",
            "Total Quantity Sold",
            "Date Range",
        ],
        "Value": [
            total,
            verified,
            (total - verified) / total if total else 0,
            df["Transaction ID"].nunique(),
            df["Customer ID"].nunique(),
            df["Total Spent"].mean(),
            df["Quantity"].sum(),
            f"{df['Transaction Date'].min():%Y-%m-%d} to {df['Transaction Date'].max():%Y-%m-%d}",
        ],
    })

    def group_summary(by):
        return (
            df.groupby(by)
            .agg(
                Revenue=("Total Spent", "sum"),
                Transactions=("Transaction ID", "nunique"),
                Quantity_Sold=("Quantity", "sum"),
                Average_Transaction=("Total Spent", "mean"),
            )
            .sort_values("Revenue", ascending=False)
            .reset_index()
        )

    monthly = (
        df.dropna(subset=["Transaction Date"])
        .assign(Year_Month=lambda d: d["Transaction Date"].dt.to_period("M").astype(str))
        .groupby("Year_Month")
        .agg(Revenue=("Total Spent", "sum"), Transactions=("Transaction ID", "nunique"), Quantity_Sold=("Quantity", "sum"))
        .reset_index()
    )

    customer = (
        df.groupby("Customer ID")
        .agg(Revenue=("Total Spent", "sum"), Transactions=("Transaction ID", "nunique"), Average_Transaction=("Total Spent", "mean"), Last_Purchase=("Transaction Date", "max"))
        .sort_values("Revenue", ascending=False)
        .reset_index()
    )
    customer["Last_Purchase"] = customer["Last_Purchase"].dt.strftime("%Y-%m-%d")

    return {
        "Executive Summary": kpi,
        "Category Analysis": group_summary("Category"),
        "Payment Analysis": group_summary("Payment Method"),
        "Location Analysis": group_summary("Location"),
        "Customer Analysis": customer,
        "Monthly Trend": monthly,
    }


def style_workbook(path: Path):
    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = "Arial"

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = fill
            cell.font = Font(name=font, color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = True
        for idx, col in enumerate(ws.columns, start=1):
            values = [c.value for c in col[:500] if c.value is not None]
            width = max([len(str(v)) for v in values], default=10) + 2
            ws.column_dimensions[get_column_letter(idx)].width = min(width, 40)
            header = str(ws.cell(1, idx).value)
            for c in col[1:]:
                c.font = Font(name=font)
                if header in ("Revenue", "Average_Transaction", "Total Spent", "Price Per Unit"):
                    c.number_format = "#,##0.00"
                elif header == "Transaction Date":
                    c.number_format = "yyyy-mm-dd"

    ws = wb["Executive Summary"]
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 28
    for r in range(2, ws.max_row + 1):
        label = str(ws.cell(r, 1).value)
        if "Share" in label:
            ws.cell(r, 2).number_format = "0.0%"
        elif "Revenue" in label or "Average" in label:
            ws.cell(r, 2).number_format = "#,##0.00"
        elif isinstance(ws.cell(r, 2).value, (int, float)):
            ws.cell(r, 2).number_format = "#,##0"

    def add_bar(ws_name, title, anchor):
        ws = wb[ws_name]
        ch = BarChart()
        ch.title, ch.y_axis.title, ch.x_axis.title = title, "Revenue", ws["A1"].value
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        ch.legend = None
        ch.width, ch.height = 18, 9
        ws.add_chart(ch, anchor)

    add_bar("Category Analysis", "Revenue by Category", "G2")
    add_bar("Location Analysis", "Revenue by Location", "G2")

    ws = wb["Payment Analysis"]
    pie = PieChart()
    pie.title = "Revenue Share by Payment Method"
    pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(pie, "G2")

    ws = wb["Monthly Trend"]
    line = LineChart()
    line.title, line.y_axis.title, line.x_axis.title = "Monthly Revenue Trend", "Revenue", "Month"
    line.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    line.legend = None
    line.width, line.height = 22, 9
    ws.add_chart(line, "G2")

    wb.save(path)


def build_report(path: Path, sheets: dict, validation: pd.DataFrame, actions, clean: pd.DataFrame):
    actions_df = pd.DataFrame(actions, columns=["Cleaning Step", "Rows Affected", "Note"])
    audit_flags = clean[["Transaction ID", "Price_Recovered", "Item_Recovered", "Quantity_Imputed", "Total_Recalculated", "Is_Estimated"]].copy()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        sheets["Executive Summary"].to_excel(writer, sheet_name="Executive Summary", index=False)
        validation.to_excel(writer, sheet_name="Data Quality", index=False)
        actions_df.to_excel(writer, sheet_name="Cleaning Actions", index=False)
        for name in ["Category Analysis", "Payment Analysis", "Location Analysis", "Customer Analysis", "Monthly Trend"]:
            sheets[name].to_excel(writer, sheet_name=name, index=False)
        clean[MODEL_COLS].to_excel(writer, sheet_name="Clean Data", index=False)
        audit_flags.to_excel(writer, sheet_name="Audit Flags", index=False)
    style_workbook(path)


def main():
    parser = argparse.ArgumentParser(description="Retail data cleaning and reporting automation")
    parser.add_argument("--input", default="retail_store_sales.csv")
    parser.add_argument("--outdir", default=".")
    args = parser.parse_args()

    out = Path(args.outdir)
    out.mkdir(parents=True, exist_ok=True)
    clean_file = out / "Clean_Retail_Store_Sales.csv"
    report_file = out / "Retail_Data_Automated_Report.xlsx"
    audit_file = out / "Cleaning_Audit_Flags.csv"

    raw = pd.read_csv(args.input)
    log.info("Loaded %d rows x %d columns", len(raw), len(raw.columns))

    clean, actions = clean_data(raw)
    validation = validate(raw, clean)
    sheets = analyse(clean)

    # Export only the 11 business columns as the final clean dataset.
    clean[MODEL_COLS].to_csv(clean_file, index=False)

    # Keep audit/provenance flags separately.
    clean[["Transaction ID", "Price_Recovered", "Item_Recovered", "Quantity_Imputed", "Total_Recalculated", "Is_Estimated"]].to_csv(audit_file, index=False)

    build_report(report_file, sheets, validation, actions, clean)

    print("\nAUTOMATION COMPLETED")
    print("Clean dataset:", clean_file)
    print("Audit flags  :", audit_file)
    print("Excel report :", report_file)


if __name__ == "__main__":
    main()
